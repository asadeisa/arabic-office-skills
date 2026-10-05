import zipfile

from arabic_pptx import ArabicPptx, audit_pptx, fix_pptx, rtl_positions
from pptx import Presentation


def slide_xml(path, n):
    with zipfile.ZipFile(path) as z:
        return z.read(f"ppt/slides/slide{n}.xml").decode("utf8")


def build(tmp_path):
    deck = ArabicPptx(transition="fade")
    deck.title_slide("عنوان العرض", "سطر فرعي")
    deck.bullets_slide("المحاور", ["نقطة مع PostgreSQL", "نقطة مع Nuxt 3"], reveal=True)
    deck.steps_slide("المراحل", ["جمع", "تحليل", "تنبيه"], reveal=True)
    deck.table_slide("جدول", ["الطبقة", "التقنية"], [["الواجهة", "Nuxt 3"]])
    out = tmp_path / "deck.pptx"
    deck.save(out)
    return out


def test_builder_output_passes_audit(tmp_path):
    assert audit_pptx(build(tmp_path)) == []


def test_transition_before_timing(tmp_path):
    x = slide_xml(build(tmp_path), 2)
    assert '<p:transition spd="med"><p:fade/></p:transition>' in x
    assert x.index("<p:transition") < x.index("<p:timing")


def test_bullet_reveal_is_paragraph_build(tmp_path):
    x = slide_xml(build(tmp_path), 2)
    assert x.count('nodeType="clickEffect"') == 2
    assert 'build="p"' in x and '<p:pRg st="1" end="1"/>' in x


def test_real_bullets(tmp_path):
    assert "<a:buChar" in slide_xml(build(tmp_path), 2)


def test_rtl_positions_first_item_on_the_right():
    xs = rtl_positions(3, 0, 300, 80, 10)
    assert xs[0] > xs[1] > xs[2]


def test_fix_pptx_clears_audit(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "عنوان مع Nuxt 3"
    s.placeholders[1].text_frame.text = "نص مع PostgreSQL و FastAPI"
    src, dst = tmp_path / "naive.pptx", tmp_path / "fixed.pptx"
    prs.save(src)
    assert audit_pptx(src)
    fix_pptx(src, dst)
    assert audit_pptx(dst) == []
