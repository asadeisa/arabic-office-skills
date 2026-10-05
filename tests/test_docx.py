import zipfile

from arabic_docx import ArabicDocx, audit_docx, fix_docx
from docx import Document
from docx.shared import Pt


def xml(path, part="word/document.xml"):
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf8")


def build(tmp_path):
    doc = ArabicDocx()
    doc.toc()
    doc.heading("المقدمة", page_break=True)
    doc.para("نص عربي عريض مع PostgreSQL و C#.", bold=True)
    doc.bullets(["نقطة أولى", "نقطة ثانية"])
    doc.table(["البند", "التفصيل"], [["قاعدة البيانات", "PostgreSQL"]], [5, 11],
              caption="مكونات النظام")
    doc.page_numbers(skip_first=True, start=0)
    out = tmp_path / "built.docx"
    doc.save(out)
    return out


def test_builder_output_passes_audit(tmp_path):
    assert audit_docx(build(tmp_path)) == []


def test_bold_arabic_sets_complex_script_bold(tmp_path):
    body = xml(build(tmp_path))
    assert "<w:b/><w:bCs/>" in body


def test_table_geometry_written_everywhere(tmp_path):
    body = xml(build(tmp_path))
    assert '<w:gridCol w:w="2835"/>' in body           # 5 cm
    assert '<w:tblLayout w:type="fixed"/>' in body
    assert "<w:tblHeader/>" in body and "<w:bidiVisual/>" in body


def test_toc_has_cached_entries(tmp_path):
    body = xml(build(tmp_path))
    assert "TOC \\o" in body and "المقدمة" in body.split("TOC \\o")[1]


def test_rpr_children_in_schema_order(tmp_path):
    body = xml(build(tmp_path))
    for run_props in body.split("<w:rPr>")[1:]:
        props = run_props.split("</w:rPr>")[0]
        if "<w:rFonts" in props and "<w:sz " in props:
            assert props.index("<w:rFonts") < props.index("<w:sz ")


def test_fix_docx_clears_audit(tmp_path):
    d = Document()
    p = d.add_paragraph()
    r = p.add_run("نص عربي عريض مع PostgreSQL و FastAPI")
    r.bold, r.font.size = True, Pt(14)
    src, dst = tmp_path / "naive.docx", tmp_path / "fixed.docx"
    d.save(src)
    assert audit_docx(src)
    fix_docx(src, dst)
    assert audit_docx(dst) == []
