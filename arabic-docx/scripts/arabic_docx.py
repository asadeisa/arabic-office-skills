"""Arabic (RTL) Word document builder on top of python-docx.

The critical difference from PDF generation
-------------------------------------------
Do NOT reshape text with `arabic_reshaper` here. reportlab and Pillow have no
text engine, so they must be handed final visual-order glyphs. Word has a full
bidirectional text engine and shapes at render time — feeding it pre-shaped
presentation forms makes it shape them a second time, producing broken glyphs
and text that cannot be searched, copied, or spell-checked.

Word needs **logical-order, unshaped Unicode** plus the right direction flags.
python-docx exposes almost none of those flags, so they are written directly
into the underlying XML here.

What actually has to be set (each one matters independently):

| Where | Element | Without it |
|---|---|---|
| settings.xml | `<w:themeFontLang w:bidi="ar-SA"/>` | Word never engages its RTL layout engine |
| styles.xml docDefaults | `<w:lang w:bidi="ar-SA"/>`, `<w:bidi/>`, `<w:jc w:val="start"/>` | New paragraphs revert to LTR |
| each `<w:p>` | `<w:bidi/>` + `<w:jc w:val="start"/>` | Paragraph indents and alignment mirror wrongly |
| each `<w:r>` | `<w:rtl/>` | Runs reorder incorrectly in mixed text |
| each `<w:tbl>` | `<w:bidiVisual/>` | Column order flips — first column lands on the left |
| each `<w:sectPr>` | `<w:bidi/>` | Margins, columns and page furniture stay LTR |

Use `w:jc w:val="start"` rather than `right`: `start` is direction-relative and
stays correct if a paragraph turns out to be LTR, whereas hardcoded `right`
breaks mixed documents.

Fonts occupy three independent slots per run. Arabic is drawn from the
complex-script slot, so setting only `w:ascii`/`w:hAnsi` silently leaves Arabic
on a default face at the wrong size:

    <w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>   + <w:szCs w:val="24"/>

Mixed Arabic/Latin text is split automatically into separate runs, with `w:rtl`
set only on the Arabic ones. A single run holding both scripts is the most
common defect in generated Arabic documents.

Usage
-----
    from arabic_docx import ArabicDocx, docx_to_pdf

    doc = ArabicDocx(cs_font="Arial", size=12)
    doc.title("عنوان المستند")
    doc.heading("أولاً — المقدمة")
    doc.para("نص عربي مع مصطلحات مثل PostgreSQL و FastAPI.")
    doc.table(["البند", "التفصيل"], [["قاعدة البيانات", "PostgreSQL"]], [5, 11])
    doc.bullets(["نقطة أولى", "نقطة ثانية"])
    doc.save("out.docx")
"""

from __future__ import annotations

import os
from copy import deepcopy
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# Arabic-script ranges: Arabic, Supplement, Extended-A, Presentation Forms A/B.
ARABIC_RE = re.compile(
    r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]"
)
LATIN_RE = re.compile(r"[A-Za-z0-9]")
LATIN_LETTER_RE = re.compile(r"[A-Za-z]")
ARABIC_INDIC = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")

# Mirrored characters: rendered flipped when they land in a right-to-left run.
# Both halves of a pair therefore have to share one run — see segment().
BRACKETS = {"(": ")", "[": "]", "{": "}", "<": ">", "«": "»", "‹": "›",
            "〈": "〉", "⟨": "⟩"}
CLOSERS = {close: open_ for open_, close in BRACKETS.items()}


# --------------------------------------------------------------------------
# Low-level XML helpers
# --------------------------------------------------------------------------

# OOXML property containers are xsd:sequence — child order is part of the
# schema. Word tolerates most misordering, but validators, LibreOffice and
# other consumers do not always, so every property is inserted at its schema
# position rather than appended. Orders from ECMA-376 Part 1, §17.
_ORDER = {
    "w:rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
              "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
              "noProof", "snapToGrid", "vanish", "webHidden", "color",
              "spacing", "w", "kern", "position", "sz", "szCs", "highlight",
              "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs",
              "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
    "w:pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
              "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
              "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
              "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
              "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap",
              "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "w:tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual",
                "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
                "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout",
                "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
    "w:sectPr": ["headerReference", "footerReference", "footnotePr",
                 "endnotePr", "type", "pgSz", "pgMar", "paperSrc", "pgBorders",
                 "lnNumType", "pgNumType", "cols", "formProt", "vAlign",
                 "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter",
                 "docGrid", "printerSettings", "sectPrChange"],
    "w:tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders",
               "shd", "noWrap", "tcMar", "textDirection", "tcFitText",
               "vAlign", "hideMark"],
    "w:trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore",
               "wAfter", "cantSplit", "trHeight", "tblHeader",
               "tblCellSpacing", "jc", "hidden"],
}
# settings.xml is long; only the elements this module writes need anchoring.
_SETTINGS_AFTER_UPDATEFIELDS = ["hdrShapeDefaults", "footnotePr", "endnotePr",
                                "compat", "docVars", "rsids", "mathPr",
                                "attachedSchema", "themeFontLang",
                                "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                                "doNotAutoCompressPictures", "forceUpgrade",
                                "captions", "readModeInkLockDown",
                                "smartTagType", "schemaLibrary",
                                "shapeDefaults", "doNotEmbedSmartTags",
                                "decimalSymbol", "listSeparator"]
_SETTINGS_AFTER_THEMEFONTLANG = _SETTINGS_AFTER_UPDATEFIELDS[
    _SETTINGS_AFTER_UPDATEFIELDS.index("themeFontLang") + 1:]

_THEME_ATTRS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")


def _local(el):
    return el.tag.rsplit("}", 1)[-1]


def _insert_ordered(parent, child, successors):
    """Insert `child` before the first existing sibling named in `successors`."""
    for sib in parent:
        if _local(sib) in successors:
            sib.addprevious(child)
            return child
    parent.append(child)
    return child


def _sub(parent, tag):
    """Return child `tag`, creating it at its schema position if absent."""
    el = parent.find(qn(tag))
    if el is not None:
        return el
    el = parent.makeelement(qn(tag), {})
    order = _ORDER.get(parent.tag.replace(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}", "w:"))
    name = tag.split(":", 1)[1]
    if order and name in order:
        return _insert_ordered(parent, el, set(order[order.index(name) + 1:]))
    if tag == "w:updateFields":
        return _insert_ordered(parent, el, set(_SETTINGS_AFTER_UPDATEFIELDS))
    if tag == "w:themeFontLang":
        return _insert_ordered(parent, el, set(_SETTINGS_AFTER_THEMEFONTLANG))
    parent.append(el)
    return el


def _flag(parent, tag, val=None):
    """Set a boolean-ish OOXML flag element such as <w:bidi/> or <w:rtl/>."""
    el = _sub(parent, tag)
    if val is not None:
        el.set(qn("w:val"), val)
    return el


def _unflag(parent, tag):
    el = parent.find(qn(tag))
    if el is not None:
        parent.remove(el)


def _is_on(el):
    """OOXML on/off: present with no val, or val in true/1/on, means on."""
    if el is None:
        return False
    return el.get(qn("w:val"), "true") not in ("0", "false", "off")


def _set_fonts(rpr, latin_font, cs_font, force=True):
    """Write explicit font slots and drop theme references.

    A theme attribute (`w:asciiTheme`, `w:cstheme` …) takes precedence over the
    explicit name in the same element, so writing `w:cs="Arial"` next to an
    existing `w:cstheme="minorBidi"` changes nothing: Arabic keeps rendering in
    the theme's complex-script face. The theme attributes have to go.
    """
    fonts = _sub(rpr, "w:rFonts")
    had_theme = any(fonts.get(qn(f"w:{a}")) is not None for a in _THEME_ATTRS)
    for a in _THEME_ATTRS:
        fonts.attrib.pop(qn(f"w:{a}"), None)
    if force or had_theme or fonts.get(qn("w:ascii")) is None:
        fonts.set(qn("w:ascii"), latin_font)
        fonts.set(qn("w:hAnsi"), latin_font)
    if force or had_theme or fonts.get(qn("w:cs")) is None:
        fonts.set(qn("w:cs"), cs_font)
    return fonts


def _mirror_cs(rpr):
    """Copy Latin-slot bold/italic/size into the complex-script slot.

    Arabic is formatted from `w:bCs`, `w:iCs` and `w:szCs`, never from
    `w:b`, `w:i` and `w:sz`. python-docx (`run.bold`, `style.font.size`) only
    writes the Latin half, so Arabic silently stays regular weight and default
    size while the Latin words next to it come out bold and large.
    """
    for latin, cs in (("w:b", "w:bCs"), ("w:i", "w:iCs")):
        el = rpr.find(qn(latin))
        if el is not None and rpr.find(qn(cs)) is None:
            new = _flag(rpr, cs)
            if el.get(qn("w:val")) is not None:
                new.set(qn("w:val"), el.get(qn("w:val")))
    sz = rpr.find(qn("w:sz"))
    if sz is not None and sz.get(qn("w:val")):
        _flag(rpr, "w:szCs", sz.get(qn("w:val")))


def normalize_styles_rtl(doc, cs_font="Arial", latin_font="Arial"):
    """Make every style in a document safe for Arabic.

    Templates — Word's default one included — are authored for Latin text and
    carry three defects that per-run formatting cannot reach, because they
    surface in text Word generates itself (the table of contents, captions,
    headers) and in text the user types later:

      - theme font references, so Arabic headings render in Calibri Light;
      - `w:sz` without a matching `w:szCs` and `w:b` without `w:bCs`, so
        Arabic headings are smaller and lighter than designed;
      - `<w:bidi w:val="0"/>` on a style, which overrides the RTL default.
    """
    styles = doc.styles.element
    for rpr in styles.iter(qn("w:rPr")):
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is not None:
            _set_fonts(rpr, latin_font, cs_font, force=False)
        _mirror_cs(rpr)
    for ppr in styles.iter(qn("w:pPr")):
        bidi = ppr.find(qn("w:bidi"))
        if bidi is not None and not _is_on(bidi):
            ppr.remove(bidi)


def _add_field(paragraph, instr, cached=""):
    """Append a complex field (TOC, PAGE, SEQ …) with a cached result.

    The cached text is what every viewer shows until fields are updated, and
    LibreOffice never updates them on conversion — so it must be meaningful.
    Returns the run holding the cached text.
    """
    def fld(kind):
        r = paragraph.add_run()
        el = r._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): kind})
        r._r.append(el)
        return r

    fld("begin")
    r = paragraph.add_run()
    it = r._r.makeelement(qn("w:instrText"), {})
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = f" {instr.strip()} "
    r._r.append(it)
    fld("separate")
    result = paragraph.add_run(cached)
    fld("end")
    return result


# Neutral characters that belong to a Latin token when written flush against it:
# `+963`, `.NET`, `$5`, `@user` on the leading side; `C#`, `C++`, `50%` on the
# trailing side. Left as neutrals they resolve Arabic, land in the RTL run and
# are drawn on the far side of the token — `C#` renders as `#C`.
# `%` is deliberately absent: Arabic convention puts it on the left of the
# number (٪٥٠ / %50), which is what the bidi algorithm already does.
AFFIX_BEFORE = set("+-.#@$/~\\")
AFFIX_AFTER = set("+#")
LRE, PDF = "\u202a", "\u202c"   # LEFT-TO-RIGHT EMBEDDING … POP DIRECTIONAL FORMATTING


def _guard(chunk):
    """Fence a Latin run that starts or ends with a flush affix.

    Word and PowerPoint decide a neutral's direction from its run, so `C#` in
    its own `en-US` run is already right there. LibreOffice, PDF viewers and
    browsers ignore run boundaries and run the plain Unicode bidi algorithm
    across the paragraph, where `C#` at the edge of a Latin island renders
    `#C` and `.NET` renders `NET.`. An LRE…PDF pair around the token fixes it
    in both families. Measured on Word, PowerPoint and LibreOffice:

        plain run       Word ✓  PowerPoint ✓  LibreOffice ✗
        LRM on the edge Word ✗  PowerPoint ✓  LibreOffice ✓
        LRE … PDF       Word ✓  PowerPoint ✓  LibreOffice ✓
        LRI … PDI       correct, but Word draws the marks as visible boxes

    Only tokens that need it are wrapped; the marks are invisible.
    """
    if chunk and (chunk[0] in AFFIX_BEFORE or chunk[-1] in AFFIX_AFTER):
        return LRE + chunk + PDF
    return chunk


def _attach_affixes(text, kinds, resolved):
    """Pull flush prefixes/suffixes into the Latin run they are attached to.

    Sentence punctuation is deliberately not in the sets: the full stop in
    `مع Nuxt 3.` ends the Arabic sentence and must stay on the Arabic side.
    """
    n = len(text)
    for i in range(n):
        if kinds[i] != "L":
            continue
        j = i - 1
        while j >= 0 and kinds[j] == "N" and text[j] in AFFIX_BEFORE:
            j -= 1
        if j < i - 1 and (j < 0 or text[j].isspace() or kinds[j] == "L"
                          or text[j] in BRACKETS):
            resolved[j + 1:i] = ["L"] * (i - 1 - j)
        k = i + 1
        while k < n and kinds[k] == "N" and text[k] in AFFIX_AFTER:
            k += 1
        if k > i + 1:
            resolved[i + 1:k] = ["L"] * (k - i - 1)


def _bracket_pairs(text):
    """Return (open_index, close_index) for every matched mirrored pair."""
    stack, pairs = [], []
    for i, ch in enumerate(text):
        if ch in BRACKETS:
            stack.append((ch, i))
        elif ch in CLOSERS:
            for k in range(len(stack) - 1, -1, -1):
                if stack[k][0] == CLOSERS[ch]:
                    pairs.append((stack[k][1], i))
                    del stack[k:]
                    break
    return pairs


def segment(text):
    """Split text into (chunk, is_arabic) pieces.

    Neutral characters — spaces, punctuation, brackets — carry no direction of
    their own, so which run they land in decides where they appear. In an RTL
    paragraph a neutral belongs with the Arabic side unless it sits *between*
    two Latin tokens. Attaching a boundary space to the Latin run instead makes
    it collapse at the direction switch, which is what produces the classic
    `PostgreSQLو` with the space visually swallowed.

    Consequences of the rule:
      "Vue.js 3 مع Nuxt 3."  → ["Vue.js 3"] [" مع "] ["Nuxt 3"] ["."]
      "(20 فأكثر)"           → ["("] ["20"] [" فأكثر)"]

    That rule alone splits bracket pairs, which is a second, worse defect. In
    `DECIMAL(5,2)` the inner `(` sits between two Latin tokens and stays Latin,
    while the trailing `)` has Arabic after it and goes to the Arabic run —
    where Word mirrors it into `(` and moves it to the far edge of the Latin
    island, rendering `(DECIMAL(5,2`. So a matched pair is forced into one run
    whenever either half resolved Latin.

    `(API)` and `(1.25)` need no such repair and must not get it: there both
    brackets resolve Arabic, both are mirrored, and their positions swap, which
    cancels out. Only pairs that are actually split are touched. The span is
    left alone if it contains Arabic — that means an unpaired `<`/`>` used as a
    comparison operator, not a bracket.
    """
    if not text:
        return []

    def classify(ch):
        if ARABIC_RE.match(ch):
            return "A"
        return "L" if LATIN_RE.match(ch) else "N"

    kinds = [classify(ch) for ch in text]
    n = len(kinds)

    # Nearest strong kind on each side, so neutrals can be resolved in one pass.
    prev_strong, last = [None] * n, None
    for i, k in enumerate(kinds):
        prev_strong[i] = last
        if k != "N":
            last = k
    next_strong, nxt = [None] * n, None
    for i in range(n - 1, -1, -1):
        next_strong[i] = nxt
        if kinds[i] != "N":
            nxt = kinds[i]

    resolved = [
        k if k != "N"
        else ("L" if prev_strong[i] == "L" and next_strong[i] == "L" else "A")
        for i, k in enumerate(kinds)
    ]
    _attach_affixes(text, kinds, resolved)

    for i, j in _bracket_pairs(text):
        if (resolved[i] == "L" or resolved[j] == "L") and "A" not in kinds[i:j + 1]:
            resolved[i:j + 1] = ["L"] * (j - i + 1)

    out, buf, cur = [], text[0], resolved[0]
    for ch, k in zip(text[1:], resolved[1:]):
        if k == cur:
            buf += ch
        else:
            out.append((buf, cur == "A"))
            buf, cur = ch, k
    out.append((buf, cur == "A"))
    return [(t, a) for t, a in out if t]


def arabic_digits(text, keep_latin_context=False):
    """Convert Western digits to Arabic-Indic (٠–٩).

    With `keep_latin_context=True`, a number stays Western when the nearest
    word before or after it is Latin — `Nuxt 3`, `Python 3.12`, `ISO 9001` —
    and only numbers in Arabic context are converted. Use one digit system
    per document; mixing ٣ and 3 in running text reads as an error.
    """
    if not keep_latin_context:
        return text.translate(ARABIC_INDIC)

    def repl(m):
        before = re.search(r"(\S+)\s*$", text[:m.start()])
        after = re.match(r"\s*(\S+)", text[m.end():])
        for word in (before, after):
            if word and LATIN_LETTER_RE.search(word.group(1)) \
                    and not ARABIC_RE.search(word.group(1)):
                return m.group(0)
        if (m.start() and text[m.start() - 1].isalpha() and not ARABIC_RE.match(text[m.start() - 1])):
            return m.group(0)
        return m.group(0).translate(ARABIC_INDIC)

    return re.sub(r"\d+(?:[.,]\d+)*", repl, text)


# --------------------------------------------------------------------------
# Builder
# --------------------------------------------------------------------------

CAPTION_LABELS = {"figure": "الشكل", "table": "الجدول"}


class ArabicDocx:
    """Build a right-to-left Word document.

    cs_font    font used for the Arabic (complex-script) slot
    latin_font font used for Latin text
    size       base point size, applied to both slots
    template   optional .docx/.dotx whose styles, margins and headers to keep;
               see also `ArabicDocx.from_template`
    """

    def __init__(self, cs_font="Arial", latin_font="Arial", size=12, template=None):
        self.doc = Document(template) if template else Document()
        self.cs_font, self.latin_font, self.size = cs_font, latin_font, size
        self._anchor = None          # set by insert_before()
        self._headings = []          # (level, text) — fills the TOC fallback
        self._toc = None
        self._counters = {}
        self._patch_settings()
        self._patch_styles()
        normalize_styles_rtl(self.doc, cs_font, latin_font)

    @classmethod
    def from_template(cls, path, clear_body=True, clear_headers=False, **kw):
        """Start from an existing document or template instead of a blank one.

        Keeps styles, page size, margins and section settings. `clear_body`
        empties the body but keeps the final `w:sectPr`, which is where the
        page setup lives — deleting it silently resets the document to Letter
        with default margins. `clear_headers` empties every header and footer
        part at the XML level: `Paragraph.clear()` keeps legacy fields such as
        an old PAGE field, which then reappear next to the new ones.
        """
        self = cls(template=path, **kw)
        if clear_body:
            body = self.doc.element.body
            for child in list(body):
                if child.tag != qn("w:sectPr"):
                    body.remove(child)
        if clear_headers:
            for section in self.doc.sections:
                for part in (section.header, section.footer,
                             section.first_page_header, section.first_page_footer,
                             section.even_page_header, section.even_page_footer):
                    if part.is_linked_to_previous:
                        continue
                    el = part._element
                    for child in list(el):
                        el.remove(child)
                    el.append(el.makeelement(qn("w:p"), {}))
        return self

    # -- document-wide switches -------------------------------------------

    def _patch_settings(self):
        """themeFontLang is the switch Word actually checks for RTL layout."""
        settings = self.doc.settings.element
        tfl = _sub(settings, "w:themeFontLang")
        tfl.set(qn("w:val"), tfl.get(qn("w:val")) or "en-US")
        tfl.set(qn("w:bidi"), "ar-SA")

    def _patch_styles(self):
        styles = self.doc.styles.element
        defaults = styles.find(qn("w:docDefaults"))
        if defaults is None:
            defaults = styles.makeelement(qn("w:docDefaults"), {})
            styles.insert(0, defaults)
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is None:
            rpr_default = defaults.makeelement(qn("w:rPrDefault"), {})
            defaults.insert(0, rpr_default)
        rpr = _sub(rpr_default, "w:rPr")
        _set_fonts(rpr, self.latin_font, self.cs_font)
        lang = _sub(rpr, "w:lang")
        lang.set(qn("w:val"), lang.get(qn("w:val")) or "en-US")
        lang.set(qn("w:bidi"), "ar-SA")

        ppr_default = defaults.find(qn("w:pPrDefault"))
        if ppr_default is None:
            ppr_default = defaults.makeelement(qn("w:pPrDefault"), {})
            defaults.append(ppr_default)
        ppr = _sub(ppr_default, "w:pPr")
        _flag(ppr, "w:bidi")
        _flag(ppr, "w:jc", "start")

    def _patch_sections(self):
        for section in self.doc.sections:
            _flag(section._sectPr, "w:bidi")

    # -- placement ---------------------------------------------------------

    def _place(self, el):
        """Move a freshly appended block to the insertion anchor, if any."""
        if self._anchor is not None:
            self._anchor.addprevious(el)
        return el

    def _new_p(self, style=None):
        if style is not None and style not in [s.name for s in self.doc.styles]:
            style = None
        p = self.doc.add_paragraph(style=style)
        self._place(p._p)
        return p

    def insert_before(self, anchor_text):
        """Send every following block in front of one existing paragraph.

        For editing a document someone has already worked on — rebuilding it
        from scratch would throw their manual edits away. The anchor must match
        exactly one paragraph (exact text first, then substring); zero or
        several matches raise instead of guessing, because inserting a chapter
        in the wrong place is worse than stopping.
        """
        paras = [p for p in self.doc.paragraphs if p.text.strip()]
        hits = [p for p in paras if p.text.strip() == anchor_text.strip()]
        if not hits:
            hits = [p for p in paras if anchor_text.strip() in p.text]
        if len(hits) != 1:
            raise ValueError(f"insert_before({anchor_text!r}): expected exactly "
                             f"one matching paragraph, found {len(hits)}")
        self._anchor = hits[0]._p
        return self

    def insert_at_end(self):
        """Undo insert_before(): new blocks go to the end of the body again."""
        self._anchor = None
        return self

    # -- paragraph / run internals ----------------------------------------

    def _rtl_para(self, p, align="start"):
        """Mark a paragraph RTL and align it.

        Alignment is always written as `start`, `end`, `center` or `both`.
        Never `left`/`right`: in a bidi paragraph Word reads them relative to
        the paragraph direction (MS-OE376 §2.3.1.13 — "left is the right side
        of a right-to-left paragraph"), while other consumers read them
        physically, so the same file aligns differently in different apps.
        python-docx's WD_ALIGN_PARAGRAPH.RIGHT writes exactly that ambiguity.
        """
        align = {"right": "start", "left": "end", "centre": "center",
                 "justify": "both"}.get(align, align)
        ppr = p._p.get_or_add_pPr()
        _flag(ppr, "w:bidi")
        _flag(ppr, "w:jc", align)
        return p

    def _format_run(self, rpr, is_arabic, bold=False, italic=False, size=None):
        _set_fonts(rpr, self.latin_font, self.cs_font)
        if bold:
            _flag(rpr, "w:b")
            _flag(rpr, "w:bCs")      # Arabic bold lives here, not in w:b
        if italic:
            _flag(rpr, "w:i")
            _flag(rpr, "w:iCs")
        if size:
            half = str(round(size * 2))
            _flag(rpr, "w:sz", half)
            _flag(rpr, "w:szCs", half)   # without it Arabic ignores the size
        if is_arabic:
            _flag(rpr, "w:rtl")

    def _add_runs(self, p, text, bold=False, size=None, italic=False):
        """Split mixed text and mark only the Arabic runs as RTL."""
        pts = size or self.size
        for chunk, is_arabic in segment(text):
            run = p.add_run(chunk if is_arabic else _guard(chunk))
            self._format_run(run._r.get_or_add_rPr(), is_arabic, bold, italic, pts)
        return p

    # -- public API --------------------------------------------------------

    def title(self, text, size=None):
        p = self._new_p()
        self._rtl_para(p, align="center")
        self._add_runs(p, text, bold=True, size=size or self.size + 6)
        return self

    def subtitle(self, text):
        p = self._new_p()
        self._rtl_para(p, align="center")
        self._add_runs(p, text, size=self.size)
        return self

    def heading(self, text, size=None, level=1, page_break=False):
        """Section heading on Word's built-in `Heading N` style.

        A real heading style — not a bold paragraph — is what feeds the table
        of contents, the navigation pane and PDF bookmarks. `page_break=True`
        sets `w:pageBreakBefore` on the heading itself, which survives later
        edits; a separate page-break paragraph turns into a blank page as soon
        as the text before it grows by one line.
        """
        p = self._new_p(f"Heading {level}")
        ppr = p._p.get_or_add_pPr()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        if page_break:
            _flag(ppr, "w:pageBreakBefore")
        self._rtl_para(p)
        # Keeps the TOC working even on a template without Heading styles.
        _flag(ppr, "w:outlineLvl", str(level - 1))
        default = {1: self.size + 4, 2: self.size + 2}.get(level, self.size + 1)
        self._add_runs(p, text, bold=True, size=size or default)
        self._headings.append((level, text))
        return self

    def para(self, text, bold=False, italic=False, align="start", size=None):
        p = self._new_p()
        self._rtl_para(p, align=align)
        self._add_runs(p, text, bold=bold, italic=italic, size=size)
        return self

    def bullets(self, items, marker=None):
        """Bulleted list on the template's `List Bullet` style.

        A real list mirrors in RTL — bullet on the right, hanging indent on the
        right — and wrapped lines align under the text, not under the bullet.
        Pass a `marker` string (or use a template without `List Bullet`) to
        fall back to typed markers.
        """
        for item in items:
            p = self._new_p(None if marker else "List Bullet")
            self._rtl_para(p)
            if p.style.name != "List Bullet":
                item = f"{marker or '–'}  {item}"
            self._add_runs(p, item)
        return self

    def numbered(self, items):
        """Numbered list on the template's `List Number` style."""
        for i, item in enumerate(items, 1):
            p = self._new_p("List Number")
            self._rtl_para(p)
            if p.style.name != "List Number":
                item = f"{i}.  {item}"
            self._add_runs(p, item)
        return self

    def caption(self, text, kind="figure", label=None):
        """Numbered caption: «الشكل 3: …» / «الجدول 2: …».

        The number is a SEQ field, so Word renumbers when figures move and can
        build a list of figures from it; the cached value is filled in so the
        number shows correctly before any field update. Convention: table
        captions go above the table, figure captions below the figure.
        """
        label = label or CAPTION_LABELS.get(kind, kind)
        n = self._counters[label] = self._counters.get(label, 0) + 1
        p = self._new_p("Caption")
        self._rtl_para(p, align="center")
        p.paragraph_format.space_after = Pt(10)
        if kind == "table":
            p.paragraph_format.keep_with_next = True
        self._add_runs(p, f"{label} ", bold=True, size=self.size - 1)
        res = _add_field(p, f"SEQ {label} \\* ARABIC", str(n))
        self._format_run(res._r.get_or_add_rPr(), False, True, size=self.size - 1)
        self._add_runs(p, f": {text}", size=self.size - 1)
        return self

    def figure(self, path, caption=None, width_cm=None, alt=None):
        """Centred image, optional numbered caption below it, alt text set.

        The image paragraph is kept with its caption, so the two never split
        across a page. Images wider than the text block are scaled down to fit.
        Do not centre images vertically with empty spacer paragraphs: they
        drift onto the next page as soon as anything above them changes.
        """
        p = self._new_p()
        self._rtl_para(p, align="center")
        if caption:
            p.paragraph_format.keep_with_next = True
        run = p.add_run()
        section = self.doc.sections[-1]
        text_width = section.page_width - section.left_margin - section.right_margin
        pic = run.add_picture(str(path), width=Cm(width_cm) if width_cm else None)
        if pic.width > text_width:
            ratio = text_width / pic.width
            pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
        doc_pr = run._r.find(".//" + qn("wp:docPr"))
        if doc_pr is not None:
            description = alt or caption or ""
            doc_pr.set("descr", description)
            doc_pr.set("title", description)
        if caption:
            self.caption(caption, kind="figure")
        return self

    def table(self, header, rows, widths_cm=None, style="Table Grid",
              repeat_header=True, caption=None, header_fill=None):
        """Add an RTL table.

        `header` and each row are given in natural reading order — the first
        element is the rightmost column. `w:bidiVisual` performs the flip, so
        the column order is NOT reversed here (unlike the PDF builder, where
        reportlab has no equivalent and reversal must be done by hand).

        widths_cm     written to all three places Word reads widths from —
                      `w:tblW`, every `w:gridCol` and every `w:tcW` — with a
                      fixed layout. Setting `cell.width` alone leaves the grid
                      at its defaults and Word redistributes the columns.
        repeat_header repeats the header row on every page; rows never split.
        caption       numbered «الجدول N» caption placed above the table.
        header_fill   hex shading for the header row, e.g. "D9E2F3".
        """
        if caption:
            self.caption(caption, kind="table")
        t = self.doc.add_table(rows=0, cols=len(header))
        self._place(t._tbl)
        try:
            t.style = style
        except (KeyError, ValueError):
            self._borders(t)       # template without "Table Grid"
        tblPr = t._tbl.tblPr
        _flag(tblPr, "w:bidiVisual")

        def fill(values, is_header):
            row = t.add_row()
            trPr = row._tr.get_or_add_trPr()
            _flag(trPr, "w:cantSplit")
            if is_header and repeat_header:
                _flag(trPr, "w:tblHeader")
            for cell, value in zip(row.cells, values):
                cell.text = ""
                p = cell.paragraphs[0]
                self._rtl_para(p)
                self._add_runs(p, str(value), bold=is_header, size=self.size - 0.5)
                if is_header and header_fill:
                    shd = _sub(cell._tc.get_or_add_tcPr(), "w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), header_fill)

        fill(header, True)
        for row in rows:
            fill(row, False)

        if widths_cm:
            self._geometry(t, widths_cm)
        return self

    def _geometry(self, t, widths_cm):
        twips = [round(w * 567) for w in widths_cm]     # 1 cm = 566.9 twips
        tblPr = t._tbl.tblPr
        tblW = _sub(tblPr, "w:tblW")
        tblW.set(qn("w:w"), str(sum(twips)))
        tblW.set(qn("w:type"), "dxa")
        _sub(tblPr, "w:tblLayout").set(qn("w:type"), "fixed")
        grid = t._tbl.tblGrid
        cols = grid.findall(qn("w:gridCol"))
        for col, w in zip(cols, twips):
            col.set(qn("w:w"), str(w))
        for row in t.rows:
            for cell, w in zip(row.cells, twips):
                tcW = _sub(cell._tc.get_or_add_tcPr(), "w:tcW")
                tcW.set(qn("w:w"), str(w))
                tcW.set(qn("w:type"), "dxa")

    @staticmethod
    def _borders(t, size=4, color="808080"):
        borders = _sub(t._tbl.tblPr, "w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = _sub(borders, f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)

    def toc(self, title="المحتويات", levels="1-3"):
        """Table of contents as a real TOC field.

        Word fills in page numbers when fields are updated — `preview()` and
        `refresh_fields()` do that through Word, and the document is flagged
        so Word offers to update on open. Until then, and in LibreOffice
        (which never updates fields on conversion), the field shows the
        heading list written into its cached result at save time.
        """
        if title:
            p = self._new_p("TOC Heading")
            self._rtl_para(p)
            self._add_runs(p, title, bold=True, size=self.size + 4)
        p = self._new_p()
        self._rtl_para(p)
        self._toc = (p, f'TOC \\o "{levels}" \\h \\z \\u', levels)
        return self

    def _finish_toc(self):
        if self._toc is None:
            return
        p, instr, levels = self._toc
        self._toc = None
        top = int(levels.split("-")[-1])
        entries = [(lvl, txt) for lvl, txt in self._headings if lvl <= top] \
            or [(1, "حدِّث الفهرس لعرض المحتويات")]
        # Field codes span paragraphs: begin in the first, end in the last.
        res = _add_field(p, instr, "")
        end = res._r.getnext()                 # the fldChar end run
        res._r.getparent().remove(res._r)
        paragraphs = [p]
        for _ in entries[1:]:
            new = self.doc.add_paragraph()
            paragraphs[-1]._p.addnext(new._p)
            paragraphs.append(new)
        for para, (lvl, text) in zip(paragraphs, entries):
            style = f"toc {lvl}" if f"toc {lvl}" in [s.name for s in self.doc.styles] \
                else None
            if style:
                para.style = style
            self._rtl_para(para)
            para.paragraph_format.first_line_indent = None
            ind = _sub(para._p.get_or_add_pPr(), "w:ind")
            ind.set(qn("w:start"), str((lvl - 1) * 400))
            if para is p:
                # cached text must sit between `separate` and `end`
                for chunk, is_arabic in segment(text):
                    r = p.add_run(chunk if is_arabic else _guard(chunk))
                    self._format_run(r._r.get_or_add_rPr(), is_arabic, size=self.size)
                    end.addprevious(r._r)
            else:
                self._add_runs(para, text)
        if len(paragraphs) > 1:
            paragraphs[-1]._p.append(end)
        _flag(self.doc.settings.element, "w:updateFields", "true")

    def page_numbers(self, align="center", skip_first=False, start=None):
        """PAGE field in the footer of every section.

        skip_first  no number on the first page (a cover); sets `w:titlePg`.
        start       first page number; `skip_first=True, start=0` makes the
                    page after the cover read 1, the usual report layout.
        Previous footer content is removed at the XML level first.
        """
        for i, section in enumerate(self.doc.sections):
            footer = section.footer
            if i == 0 or not footer.is_linked_to_previous:
                footer.is_linked_to_previous = False
                el = footer._element
                for child in list(el):
                    el.remove(child)
                p = footer.add_paragraph()
                self._rtl_para(p, align=align)
                res = _add_field(p, "PAGE", "1")
                self._format_run(res._r.get_or_add_rPr(), False, size=self.size - 1)
            sectPr = section._sectPr
            if skip_first:
                section.different_first_page_header_footer = True
            if start is not None:
                _sub(sectPr, "w:pgNumType").set(qn("w:start"), str(start))
        return self

    def page_break(self):
        """Prefer heading(..., page_break=True); see heading()."""
        p = self._new_p()
        self._rtl_para(p)
        p.add_run().add_break(WD_BREAK.PAGE)
        return self

    def spacer(self, points=10):
        p = self._new_p()
        self._rtl_para(p)
        p.paragraph_format.space_after = Pt(points)
        return self

    def save(self, path):
        self._patch_sections()   # sections may be added late; patch at the end
        self._finish_toc()
        self.doc.save(str(path))
        return str(path)


# --------------------------------------------------------------------------
# Verification
# --------------------------------------------------------------------------

def find_word():
    """Locate WINWORD.EXE across Office versions. Windows only — the conversion
    below drives Word through COM, which does not exist on macOS or Linux."""
    if sys.platform != "win32":
        return None
    roots = [os.environ.get("ProgramFiles", r"C:\Program Files"),
             os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")]
    subdirs = ["root/Office16", "Office16", "root/Office15", "Office15", "Office14"]
    for root in roots:
        for sub in subdirs:
            cand = Path(root) / "Microsoft Office" / sub / "WINWORD.EXE"
            if cand.exists():
                return str(cand)
    return None


def find_soffice():
    """Locate LibreOffice — PATH first, then the usual install locations on
    Windows, macOS and Linux, since it is rarely on PATH on the first two."""
    found = shutil.which("soffice") or shutil.which("soffice.exe")
    if found:
        return found
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/bin/soffice", "/usr/local/bin/soffice", "/snap/bin/libreoffice",
    ]
    return next((c for c in candidates if Path(c).exists()), None)


NO_CONVERTER = (
    "Cannot render a preview: neither Microsoft Word nor LibreOffice was found. "
    "Install LibreOffice (https://www.libreoffice.org) — headless conversion is "
    "enough — or open the .docx on a machine that has Word. The document itself "
    "is already written; only the visual check needs a converter."
)

WORD_FAILED = (
    "Microsoft Word was found but failed to convert the document, and there is "
    "no LibreOffice install to fall back to. Word reported:\n{}"
)

# Pure ASCII, and paths arrive through the environment rather than being
# interpolated in. Windows PowerShell 5.1 reads a BOM-less script as the system
# ANSI codepage, so an Arabic path baked into the text comes out mangled, Word
# cannot find the file, and the conversion fails for every Arabic-named
# document — which is most of them. Environment variables are passed as
# Unicode, so the script never has to carry a non-ASCII character.
WORD_SCRIPT = r'''
$ErrorActionPreference = "Stop"
$src = $env:ARABIC_DOCX_SRC
$pdf = $env:ARABIC_DOCX_PDF
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {
    $doc = $word.Documents.Open($src, $false, $true)
    # Fill TOC page numbers and SEQ/PAGE fields in memory (the file is opened
    # read-only and never saved), so the preview shows what the reader will
    # see after Word's own update-fields prompt.
    foreach ($t in $doc.TablesOfContents) { $t.Update() }
    $doc.Fields.Update() | Out-Null
    $doc.SaveAs([ref]$pdf, [ref]17)    # 17 = wdFormatPDF
    $doc.Close([ref]0)
} finally {
    $word.Quit()
}
'''


def _console_text(raw):
    """Decode PowerShell output, which is console-codepage bytes, not UTF-8."""
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return repr(raw)


def _word_to_pdf(docx_path, pdf_path):
    """Drive Word through COM. Returns None on success, else the failure text."""
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False,
                                     encoding="utf-8-sig") as fh:
        fh.write(WORD_SCRIPT)
        ps = fh.name
    env = dict(os.environ, ARABIC_DOCX_SRC=str(docx_path),
               ARABIC_DOCX_PDF=str(pdf_path))
    try:
        proc = subprocess.run(["powershell", "-NoProfile", "-NonInteractive",
                               "-ExecutionPolicy", "Bypass", "-File", ps],
                              capture_output=True, env=env)
    finally:
        Path(ps).unlink(missing_ok=True)
    if proc.returncode == 0 and Path(pdf_path).exists():
        return None
    return (_console_text(proc.stderr) or _console_text(proc.stdout)
            or f"powershell exited with {proc.returncode} and wrote no PDF")


def docx_to_pdf(docx_path, pdf_path=None):
    """Convert via Word or LibreOffice so the result can be rendered and viewed.

    Arabic defects do not show up in the document's own text — the characters
    are all present and correct, only their shaping and order are wrong. Only a
    render reveals them, so convert and look before delivering.

    A Word failure is carried, not swallowed. Reporting "no converter found"
    when Word is installed and merely errored sends the caller off installing
    LibreOffice for a problem that has nothing to do with LibreOffice.
    """
    docx_path = Path(docx_path).resolve()
    pdf_path = Path(pdf_path).resolve() if pdf_path else docx_path.with_suffix(".pdf")
    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    word_error = None
    if find_word():
        word_error = _word_to_pdf(docx_path, pdf_path)
        if word_error is None:
            return str(pdf_path)

    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(WORD_FAILED.format(word_error) if word_error
                           else NO_CONVERTER)

    try:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                        "--outdir", str(pdf_path.parent), str(docx_path)],
                       check=True, capture_output=True)
    except subprocess.CalledProcessError as exc:
        detail = _console_text(exc.stderr) or _console_text(exc.stdout)
        raise RuntimeError(
            f"LibreOffice failed to convert {docx_path.name}: {detail}"
            + (f"\nWord was tried first and failed: {word_error}" if word_error else "")
        ) from exc
    # LibreOffice names the output after the *source* file, ignoring pdf_path.
    produced = pdf_path.parent / f"{docx_path.stem}.pdf"
    if produced != pdf_path and produced.exists():
        produced.replace(pdf_path)
    return str(pdf_path)


def preview(docx_path, out_dir=None, scale=2):
    """Convert to PDF and render every page to PNG. View these images."""
    import pypdfium2 as pdfium

    docx_path = Path(docx_path)
    out_dir = Path(out_dir) if out_dir else docx_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)

    pdf = pdfium.PdfDocument(docx_to_pdf(docx_path, out_dir / f"{docx_path.stem}.pdf"))
    paths = []
    for i in range(len(pdf)):
        png = out_dir / f"{docx_path.stem}_p{i + 1}.png"
        pdf[i].render(scale=scale).to_pil().save(png)
        paths.append(str(png))
    # Release the file handle: on Windows an open PdfDocument locks the PDF,
    # and the next conversion to the same path fails inside Word/PowerPoint
    # with a bare E_FAIL.
    pdf.close()
    return paths


# --------------------------------------------------------------------------
# Field refresh (Word only)
# --------------------------------------------------------------------------

REFRESH_SCRIPT = r'''
$ErrorActionPreference = "Stop"
$src = $env:ARABIC_DOCX_SRC
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {
    $doc = $word.Documents.Open($src, $false, $false)
    foreach ($t in $doc.TablesOfContents) { $t.Update() }
    $doc.Fields.Update() | Out-Null
    $doc.Save()
    $doc.Close([ref]0)
} finally {
    $word.Quit()
}
'''


def refresh_fields(docx_path):
    """Update the TOC, page and caption fields in place through Word.

    Run this on the final file before handing it over, so the table of
    contents carries real page numbers even for readers whose viewer never
    updates fields (LibreOffice, browsers, most phone apps). Returns True on
    success, False when Word is not available (non-Windows, or not installed).
    """
    if not find_word():
        return False
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False,
                                     encoding="utf-8-sig") as fh:
        fh.write(REFRESH_SCRIPT)
        ps = fh.name
    env = dict(os.environ, ARABIC_DOCX_SRC=str(Path(docx_path).resolve()))
    try:
        proc = subprocess.run(["powershell", "-NoProfile", "-NonInteractive",
                               "-ExecutionPolicy", "Bypass", "-File", ps],
                              capture_output=True, env=env)
    finally:
        Path(ps).unlink(missing_ok=True)
    if proc.returncode != 0:
        raise RuntimeError("Word failed to refresh fields: "
                           + (_console_text(proc.stderr) or _console_text(proc.stdout)))
    return True


# --------------------------------------------------------------------------
# Audit and repair of existing documents
# --------------------------------------------------------------------------

PRESENTATION_FORMS_RE = re.compile(r"[ﭐ-﷿ﹰ-﻿]")


def _style_map(doc):
    """styleId -> (pPr, rPr, basedOn) for resolving inherited properties."""
    out = {}
    for st in doc.styles.element.findall(qn("w:style")):
        based = st.find(qn("w:basedOn"))
        out[st.get(qn("w:styleId"))] = (st.find(qn("w:pPr")), st.find(qn("w:rPr")),
                                        based.get(qn("w:val")) if based is not None else None)
    return out


def _inherited(styles, style_id, kind, tag, default_el):
    """Find `tag` in a style chain (kind 0 = pPr, 1 = rPr), else docDefaults."""
    seen = set()
    while style_id and style_id not in seen and style_id in styles:
        seen.add(style_id)
        props = styles[style_id][kind]
        if props is not None and props.find(qn(tag)) is not None:
            return props.find(qn(tag))
        style_id = styles[style_id][2]
    return default_el.find(qn(tag)) if default_el is not None else None


def _all_paragraphs(doc):
    yield from doc.element.body.iter(qn("w:p"))
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header,
                     section.even_page_footer):
            if not part.is_linked_to_previous:
                yield from part._element.iter(qn("w:p"))


def _run_text(r):
    return "".join(t.text or "" for t in r.iter(qn("w:t")))


def audit_docx(path):
    """Report the Arabic defects that do not show up in a document's text.

    Returns a list of human-readable findings, empty when the document is
    clean. Each check targets a defect seen in real generated documents; the
    message says what the reader will see. Run it on generated files before
    delivery, and on files a user hands over to find out what to fix.
    """
    doc = Document(str(path))
    issues, counts, examples = [], {}, {}

    def note(key, sample):
        counts[key] = counts.get(key, 0) + 1
        if len(examples.setdefault(key, [])) < 2 and sample:
            examples[key].append(sample.strip()[:50])

    tfl = doc.settings.element.find(qn("w:themeFontLang"))
    if tfl is None or not tfl.get(qn("w:bidi")):
        issues.append("settings.xml has no <w:themeFontLang w:bidi=…>: Word may "
                      "not engage its RTL layout engine.")

    styles = _style_map(doc)
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    ppr_def = defaults.find(qn("w:pPrDefault") + "/" + qn("w:pPr")) if defaults is not None else None

    for st in doc.styles.element.findall(qn("w:style")):
        name = st.find(qn("w:name"))
        name = name.get(qn("w:val")) if name is not None else st.get(qn("w:styleId"))
        rpr = st.find(qn("w:rPr"))
        if rpr is not None:
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None and fonts.get(qn("w:cstheme")) is not None:
                note("style-theme", name)
            sz, szcs = rpr.find(qn("w:sz")), rpr.find(qn("w:szCs"))
            if sz is not None and (szcs is None or szcs.get(qn("w:val")) != sz.get(qn("w:val"))):
                note("style-szcs", name)
            if _is_on(rpr.find(qn("w:b"))) and rpr.find(qn("w:bCs")) is None:
                note("style-bcs", name)
        ppr = st.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:bidi")) is not None \
                and not _is_on(ppr.find(qn("w:bidi"))):
            note("style-bidi0", name)

    for p in _all_paragraphs(doc):
        text = "".join(_run_text(r) for r in p.iter(qn("w:r")))
        if PRESENTATION_FORMS_RE.search(text):
            note("shaped", text)
        if not ARABIC_RE.search(text):
            continue
        ppr = p.find(qn("w:pPr"))
        sid = None
        if ppr is not None and ppr.find(qn("w:pStyle")) is not None:
            sid = ppr.find(qn("w:pStyle")).get(qn("w:val"))
        bidi = ppr.find(qn("w:bidi")) if ppr is not None else None
        if bidi is None:
            bidi = _inherited(styles, sid, 0, "w:bidi", ppr_def)
        if not _is_on(bidi):
            note("no-bidi", text)
        else:
            jc = ppr.find(qn("w:jc")) if ppr is not None else None
            if jc is not None and jc.get(qn("w:val")) in ("left", "right"):
                note("jc", text)
        for r in p.iter(qn("w:r")):
            rt = _run_text(r)
            if not rt:
                continue
            has_ar = bool(ARABIC_RE.search(rt))
            rpr = r.find(qn("w:rPr"))
            if has_ar and LATIN_LETTER_RE.search(rt):
                note("mixed", rt)
            if has_ar and (rpr is None or not _is_on(rpr.find(qn("w:rtl")))):
                note("no-rtl", rt)
            if has_ar and rpr is not None:
                if _is_on(rpr.find(qn("w:b"))) and rpr.find(qn("w:bCs")) is None:
                    note("bcs", rt)
                sz, szcs = rpr.find(qn("w:sz")), rpr.find(qn("w:szCs"))
                if sz is not None and szcs is None:
                    note("szcs", rt)

    for tbl in doc.element.body.iter(qn("w:tbl")):
        if ARABIC_RE.search("".join(t.text or "" for t in tbl.iter(qn("w:t")))):
            tblPr = tbl.find(qn("w:tblPr"))
            if tblPr is None or not _is_on(tblPr.find(qn("w:bidiVisual"))):
                note("table", "")

    for i, section in enumerate(doc.sections):
        if not _is_on(section._sectPr.find(qn("w:bidi"))):
            note("section", f"section {i + 1}")

    body = doc.element.body
    toc = next((el for el in body.iter(qn("w:instrText"))
                if (el.text or "").strip().startswith("TOC")), None)
    if toc is not None:
        # Collect the cached result: text between this field's instruction and
        # its end marker. Empty means the TOC shows nothing outside Word.
        cached, after = "", False
        for el in body.iter():
            if el is toc:
                after = True
            elif after and el.tag == qn("w:fldChar") \
                    and el.get(qn("w:fldCharType")) == "end":
                break
            elif after and el.tag == qn("w:t"):
                cached += el.text or ""
        if not cached.strip():
            note("toc", "")

    messages = {
        "shaped": "paragraph(s) contain pre-shaped Arabic presentation forms — "
                  "arabic_reshaper was run on Word text; it renders broken and "
                  "cannot be searched. Write plain logical Unicode instead",
        "no-bidi": "Arabic paragraph(s) without <w:bidi/>: they start at the "
                   "left margin and mixed text reorders",
        "jc": "RTL paragraph(s) aligned with jc=left/right: Word reads these "
              "relative to direction, other apps physically — use start/end",
        "mixed": "run(s) mixing Arabic and Latin letters: spaces collapse and "
                 "Latin terms land on the wrong side — split by script",
        "no-rtl": "Arabic run(s) without <w:rtl/>",
        "bcs": "bold Arabic run(s) without <w:bCs/>: Arabic stays regular weight",
        "szcs": "Arabic run(s) with <w:sz> but no <w:szCs>: Arabic ignores the size",
        "table": "table(s) with Arabic text but no <w:bidiVisual/>: first column "
                 "renders on the left",
        "section": "section(s) without <w:bidi/>: margins and page furniture stay LTR",
        "style-theme": "style(s) using a theme complex-script font (w:cstheme): "
                       "Arabic in TOC, captions and headings falls back to the theme face",
        "style-szcs": "style(s) whose w:szCs differs from w:sz: Arabic in that "
                      "style renders at a different size than Latin",
        "style-bcs": "bold style(s) without w:bCs: Arabic in that style is not bold",
        "style-bidi0": "style(s) with <w:bidi w:val=\"0\"/>, overriding RTL",
        "toc": "table of contents field has no cached entries: it shows empty "
               "outside Word — run refresh_fields() through Word",
    }
    for key, n in counts.items():
        ex = examples.get(key) or []
        issues.append(f"{n} × {messages[key]}" + (f" — e.g. {ex}" if ex else ""))
    return issues


def _split_mixed_run(r):
    """Split one run into script-homogeneous runs, cloning its formatting.

    Only plain text runs (rPr + w:t) are split; runs carrying tabs, breaks,
    fields or drawings are left as they are and still show up in the audit.
    """
    children = [c for c in r if c.tag != qn("w:rPr")]
    if len(children) != 1 or children[0].tag != qn("w:t"):
        return [r]
    pieces = segment(children[0].text or "")
    if len(pieces) < 2:
        return [r]
    out = []
    for chunk, _ in pieces:
        new = deepcopy(r)
        t = new.find(qn("w:t"))
        t.text = chunk
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.addprevious(new)
        out.append(new)
    r.getparent().remove(r)
    return out


def fix_docx(src, dst, cs_font=None, tables=False):
    """Retrofit RTL onto an existing Word document; returns `dst`.

    For files produced elsewhere — a template, an export, a document typed in
    an LTR setup. Everything is done in place on the original structure:

      - document switches (`themeFontLang`, docDefaults) and every style
        normalised as in `normalize_styles_rtl`;
      - every paragraph containing Arabic gets `w:bidi`; left/right alignment
        becomes `start`/`end` (a paragraph that was LTR before keeps reading
        from the right margin);
      - mixed-script runs are split, Arabic runs get `w:rtl`, and bold,
        italic and size are mirrored into the complex-script slots;
      - every section gets `w:bidi`.

    `tables=True` also adds `w:bidiVisual` to tables containing Arabic. It is
    off by default because it flips the visual column order: correct for a
    table whose columns were typed in reading order, wrong for one whose
    columns were already reversed by hand to look right in LTR.
    """
    doc = Document(str(src))
    fonts = doc.styles.element.find(".//" + qn("w:rPrDefault") + "//" + qn("w:rFonts"))
    cs_font = cs_font or (fonts.get(qn("w:cs")) if fonts is not None else None) or "Arial"
    latin = (fonts.get(qn("w:ascii")) if fonts is not None else None) or "Arial"

    # Reuse the builder's document-wide patches on the loaded document.
    builder = ArabicDocx.__new__(ArabicDocx)
    builder.doc, builder.cs_font, builder.latin_font = doc, cs_font, latin
    builder._patch_settings()
    builder._patch_styles()
    normalize_styles_rtl(doc, cs_font, latin)

    for p in list(_all_paragraphs(doc)):
        text = "".join(_run_text(r) for r in p.iter(qn("w:r")))
        if not ARABIC_RE.search(text):
            continue
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = p.makeelement(qn("w:pPr"), {})
            p.insert(0, ppr)
        was_bidi = _is_on(ppr.find(qn("w:bidi")))
        _flag(ppr, "w:bidi")
        jc = ppr.find(qn("w:jc"))
        if jc is not None:
            val = jc.get(qn("w:val"))
            if not was_bidi and val in ("left", "right"):
                jc.set(qn("w:val"), "start")
            elif val == "left":
                jc.set(qn("w:val"), "start")
            elif val == "right":
                jc.set(qn("w:val"), "end")
        for r in list(p.iter(qn("w:r"))):
            for piece in _split_mixed_run(r):
                rpr = piece.find(qn("w:rPr"))
                if rpr is None:
                    rpr = piece.makeelement(qn("w:rPr"), {})
                    piece.insert(0, rpr)
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is not None and rfonts.get(qn("w:cs")) is None:
                    rfonts.set(qn("w:cs"), cs_font)
                _mirror_cs(rpr)
                if ARABIC_RE.search(_run_text(piece)):
                    _flag(rpr, "w:rtl")

    if tables:
        for tbl in doc.element.body.iter(qn("w:tbl")):
            if ARABIC_RE.search("".join(t.text or "" for t in tbl.iter(qn("w:t")))):
                _flag(tbl.find(qn("w:tblPr")), "w:bidiVisual")

    for section in doc.sections:
        _flag(section._sectPr, "w:bidi")
    doc.save(str(dst))
    return str(dst)
